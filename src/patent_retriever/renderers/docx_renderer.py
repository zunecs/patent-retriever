"""Render a PatentDocument into the format required by the Patent Translation tool.

The layout mirrors a real sample of the target document, inspected in Word:

  Times New Roman 13 pt, 1.5 line spacing, 6 pt after, justified body.
  Page: 1.2" top, 1" bottom/left/right.
  A right-aligned running page header carrying the docket and client reference.
  Body paragraph IDs are a genuine Word numbered list, not literal text.

Rendering happens in two stages. `render_blocks` produces a typed sequence of
(kind, text) pairs describing the document logically; `render_docx` maps each
kind onto Word formatting. The split keeps the document's structure assertable in
tests without parsing .docx binaries, and gives the web interface a preview that
cannot drift from the real output.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph as WordParagraph

from patent_retriever.domain.models import PatentDocument, RenderOptions

FONT_NAME = "Times New Roman"
FONT_SIZE_PT = 13
COVER_SIZE_PT = 18
LABEL_SIZE_PT = 14
LINE_SPACING = 1.5
SPACE_AFTER_PT = 6

MARGIN_TOP_IN = 1.2
MARGIN_BOTTOM_IN = 1.0
MARGIN_LEFT_IN = 1.0
MARGIN_RIGHT_IN = 1.0

# Word list geometry, read from the sample's "Adjust List Indents" dialog.
NUMBER_POSITION_IN = 0.44
TEXT_INDENT_IN = 1.04

CLAIMS_HEADING = "CLAIMS"
CLAIMS_SUBHEADING = "What is claimed:"
ABSTRACT_HEADING = "ABSTRACT"

_TWIPS_PER_INCH = 1440
_NUMBERING_ID = 900

# The abstract's closing reference is written "Fig.1." - no space after "Fig."
_FIGURE_REFERENCE = re.compile(r"fig\.?\s*(\d+[A-Za-z]?)\.?", re.IGNORECASE)


class Kind(StrEnum):
    """The logical role of a line, which decides its Word formatting."""

    COVER = "cover"  # APPLICATION / FOR / UNITED STATES LETTERS PATENT
    LABEL = "label"  # TITLE: and INVENTOR: lines
    DOC_TITLE = "doc_title"  # the title repeated above the body
    HEADING = "heading"  # BACKGROUND, SUMMARY, CLAIMS, ...
    ABSTRACT_HEADING = "abstract_heading"
    NUMBERED = "numbered"  # body paragraph carrying a list number
    CLAIMS_SUBHEADING = "claims_sub"
    CLAIM = "claim"
    BODY = "body"  # abstract text
    FIGURE = "figure"  # the closing Fig.N. line
    BLANK = "blank"
    PAGE_BREAK = "page_break"


@dataclass(frozen=True, slots=True)
class Block:
    """One rendered line and the role that determines how it is formatted."""

    kind: Kind
    text: str = ""


def normalize_figure_reference(raw: str | None) -> str | None:
    """Return the closing figure reference in the document's exact form.

    Sources write "FIG. 1", "Fig. 1." and similar; the target document uses
    "Fig.1." with no space and a trailing period.
    """
    if not raw:
        return None
    match = _FIGURE_REFERENCE.search(raw)
    if match is None:
        return None
    return f"Fig.{match.group(1)}."


def _cover_blocks(document: PatentDocument) -> list[Block]:
    """The first page: the application statement, title, and inventors.

    The sample opens with four empty paragraphs before the application statement.
    """
    inventors = ", ".join(document.inventors)
    return [
        Block(Kind.BLANK),
        Block(Kind.BLANK),
        Block(Kind.BLANK),
        Block(Kind.BLANK),
        Block(Kind.COVER, "APPLICATION"),
        Block(Kind.COVER, "FOR"),
        Block(Kind.COVER, "UNITED STATES LETTERS PATENT"),
        Block(Kind.BLANK),
        Block(Kind.BLANK),
        Block(Kind.BLANK),
        Block(Kind.LABEL, f"TITLE:\t\t{document.title}"),
        Block(Kind.BLANK),
        Block(Kind.LABEL, f"INVENTOR:\t{inventors}"),
        Block(Kind.PAGE_BREAK),
    ]


def _body_blocks(document: PatentDocument) -> list[Block]:
    """The four numbered sections. A heading is followed directly by its text."""
    blocks: list[Block] = []
    for heading, paragraphs in document.body_sections():
        if not paragraphs:
            continue
        blocks.append(Block(Kind.HEADING, heading))
        blocks.extend(Block(Kind.NUMBERED, paragraph.text) for paragraph in paragraphs)
    return blocks


def _claims_blocks(document: PatentDocument) -> list[Block]:
    """Claims are numbered as literal text, unlike the bracketed body IDs."""
    blocks = [
        Block(Kind.HEADING, CLAIMS_HEADING),
        Block(Kind.CLAIMS_SUBHEADING, CLAIMS_SUBHEADING),
        Block(Kind.BLANK),
    ]
    blocks.extend(Block(Kind.CLAIM, f"{claim.number}. {claim.text}") for claim in document.claims)
    return blocks


def _abstract_blocks(document: PatentDocument) -> list[Block]:
    if not document.abstract:
        return []
    blocks = [
        Block(Kind.ABSTRACT_HEADING, ABSTRACT_HEADING),
        Block(Kind.BODY, document.abstract),
    ]
    reference = normalize_figure_reference(document.abstract_figure_reference)
    if reference:
        blocks.append(Block(Kind.FIGURE, reference))
    return blocks


def render_blocks(document: PatentDocument, options: RenderOptions | None = None) -> list[Block]:
    """Return the document as a typed sequence of lines.

    The document is renumbered first: the output requires an unbroken paragraph
    sequence across all four body sections, which sources rarely provide.
    """
    normalized = document.renumbered()
    return [
        *_cover_blocks(normalized),
        Block(Kind.DOC_TITLE, normalized.title),
        Block(Kind.BLANK),
        *_body_blocks(normalized),
        *_claims_blocks(normalized),
        *_abstract_blocks(normalized),
    ]


def render_lines(document: PatentDocument, options: RenderOptions | None = None) -> list[str]:
    """Return the rendered text one line at a time, for previews and tests.

    Body paragraph numbers are supplied by Word's list numbering, so they do not
    appear in the text. They are added here so the preview matches what a reader
    will see in the finished document.
    """
    lines: list[str] = []
    number = 0
    for block in render_blocks(document, options):
        if block.kind is Kind.NUMBERED:
            number += 1
            lines.append(f"[{number:04d}] {block.text}")
        elif block.kind is Kind.PAGE_BREAK:
            lines.append("")
        else:
            lines.append(block.text)
    return lines


def render_text(document: PatentDocument, options: RenderOptions | None = None) -> str:
    """Return the document as plain text."""
    return "\n".join(render_lines(document, options))


def _install_numbering(word_document: docx.document.Document) -> None:
    """Define the bracketed, zero-padded list used by body paragraphs.

    Word's built-in numbering formats cannot pad to four digits, so this uses the
    `custom` format supported by ECMA-376. python-docx has no API for numbering
    definitions, so the XML is injected directly. Within numbering.xml every
    <w:abstractNum> must precede every <w:num>, which is why the element is
    inserted before the first existing <w:num> rather than appended.
    """
    numbering = word_document.part.numbering_part.element
    left = int(TEXT_INDENT_IN * _TWIPS_PER_INCH)
    hanging = int((TEXT_INDENT_IN - NUMBER_POSITION_IN) * _TWIPS_PER_INCH)

    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{_NUMBERING_ID}">'
        '<w:multiLevelType w:val="singleLevel"/>'
        '<w:lvl w:ilvl="0">'
        '<w:start w:val="1"/>'
        '<w:numFmt w:val="custom" w:format="0001, 0002, 0003, ..."/>'
        '<w:lvlText w:val="[%1]"/>'
        '<w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr>'
        "</w:lvl></w:abstractNum>"
    )
    definition = parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{_NUMBERING_ID}">'
        f'<w:abstractNumId w:val="{_NUMBERING_ID}"/></w:num>'
    )

    existing = numbering.findall(qn("w:num"))
    if existing:
        existing[0].addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(definition)


def _apply_numbering(paragraph: WordParagraph) -> None:
    """Attach a paragraph to the bracketed list defined by `_install_numbering`."""
    properties = paragraph._p.get_or_add_pPr()
    properties.append(
        parse_xml(
            f"<w:numPr {nsdecls('w')}>"
            f'<w:ilvl w:val="0"/><w:numId w:val="{_NUMBERING_ID}"/></w:numPr>'
        )
    )


def _configure_page(word_document: docx.document.Document) -> None:
    section = word_document.sections[0]
    section.top_margin = Inches(MARGIN_TOP_IN)
    section.bottom_margin = Inches(MARGIN_BOTTOM_IN)
    section.left_margin = Inches(MARGIN_LEFT_IN)
    section.right_margin = Inches(MARGIN_RIGHT_IN)

    normal = word_document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(FONT_SIZE_PT)
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(SPACE_AFTER_PT)


def _build_header(word_document: docx.document.Document, options: RenderOptions) -> None:
    """Write the running header that repeats at the top of every page.

    Right-aligned, single-spaced, and tight - it must not inherit the body's
    1.5 line spacing and 6 pt trailing space or it will push the text block down.
    """
    header = word_document.sections[0].header
    lines = [
        "PATENT APPLICATION",
        f"ATTORNEY DOCKET NO. {options.attorney_docket_number}".rstrip(),
        f"CLIENT REF. NO. {options.client_reference}".rstrip(),
    ]

    header.paragraphs[0].text = lines[0]
    for line in lines[1:]:
        header.add_paragraph(line)

    for paragraph in header.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(FONT_SIZE_PT)


def _write_block(word_document: docx.document.Document, block: Block) -> None:
    """Add one block to the document with the formatting its kind requires."""
    if block.kind is Kind.PAGE_BREAK:
        word_document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        return

    paragraph = word_document.add_paragraph()
    run = paragraph.add_run(block.text)

    if block.kind is Kind.COVER:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.font.size = Pt(COVER_SIZE_PT)
        run.font.bold = True
    elif block.kind is Kind.LABEL:
        run.font.size = Pt(LABEL_SIZE_PT)
        run.font.bold = True
    elif block.kind is Kind.DOC_TITLE or block.kind is Kind.ABSTRACT_HEADING:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.font.size = Pt(LABEL_SIZE_PT)
        run.font.bold = True
    elif block.kind is Kind.HEADING:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif block.kind is Kind.NUMBERED:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _apply_numbering(paragraph)
    elif block.kind in (Kind.CLAIM, Kind.BODY):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def render_docx(document: PatentDocument, options: RenderOptions | None = None) -> bytes:
    """Return the document as .docx bytes.

    Returning bytes rather than writing to a path keeps this function free of
    filesystem concerns, which is what lets the web interface stream the file to
    the browser without a temporary file.
    """
    options = options or RenderOptions()
    word_document = docx.Document()

    _configure_page(word_document)
    _install_numbering(word_document)
    _build_header(word_document, options)

    for block in render_blocks(document, options):
        _write_block(word_document, block)

    buffer = io.BytesIO()
    word_document.save(buffer)
    return buffer.getvalue()


def write_docx(
    document: PatentDocument,
    path: Path,
    options: RenderOptions | None = None,
) -> Path:
    """Write the rendered document to disk and return the path written."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(render_docx(document, options))
    return path
